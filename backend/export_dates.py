import docx
import openpyxl

word_doc_path = r'e:\Internship\PocketFM\US Licensing Workstream Update.docx'
doc = docx.Document(word_doc_path)
target_table = doc.tables[5]

# Extract data
extracted_data = []
for row_idx, row in enumerate(target_table.rows):
    if row_idx == 0: continue
    cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
    if len(cells) >= 7:
        title = cells[2].strip()
        date_str = cells[6].strip()
        if title and date_str:
            extracted_data.append((title, date_str))

# Create new Excel workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Correct Dates"

# Add headers
ws.append(["Title Name", "Contract Date"])

# Add data
for title, date_str in extracted_data:
    ws.append([title, date_str])

# Save
output_path = r'e:\Internship\PocketFM\Correct_Contract_Dates.xlsx'
wb.save(output_path)
print(f"Created {output_path} with {len(extracted_data)} rows.")
