import docx
import openpyxl
from datetime import datetime

# 1. Parse Word Document for Title and Date
word_doc_path = r'e:\Internship\PocketFM\US Licensing Workstream Update.docx'
doc = docx.Document(word_doc_path)

target_table = doc.tables[5]
word_data = {}

for row_idx, row in enumerate(target_table.rows):
    if row_idx == 0: continue
    cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
    if len(cells) >= 7:
        title = cells[2].strip()
        date_str = cells[6].strip()
        if title and date_str:
            word_data[title.lower()] = date_str

# 2. Parse Excel Master Sheet and Update
excel_path = r'e:\Internship\PocketFM\US_Licensing_Lifecycle_Tracker.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws_master = wb['Lifecycle Tracker - Master']

updated_count = 0

for row_idx, row in enumerate(ws_master.iter_rows(min_row=4), start=4):
    title_cell = row[4].value
    date_cell = row[13].value
    
    if not title_cell: continue
    title_excel = str(title_cell).strip()
    title_lower = title_excel.lower()
    
    if title_lower in word_data:
        word_date_str = word_data[title_lower]
        
        excel_date_str = ""
        if isinstance(date_cell, datetime):
            excel_date_str = date_cell.strftime("%Y-%m-%d")
        elif date_cell:
            excel_date_str = str(date_cell).strip().split(' ')[0]
            
        # Only overwrite if they are different
        if word_date_str != excel_date_str:
            print(f"Updating {title_excel}: {excel_date_str} -> {word_date_str}")
            row[13].value = word_date_str
            updated_count += 1

print(f"\nSuccessfully updated {updated_count} mismatched dates in the Master sheet.")
wb.save(excel_path)
print("Done!")
